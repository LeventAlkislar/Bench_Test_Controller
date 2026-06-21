# bench_test/ui/visible_xlsx_exporter.py
# -*- coding: utf-8 -*-
"""Pyqtgraph exporter for Viewer visible-range XLSX data."""

from datetime import datetime

from PyQt6.QtWidgets import QMessageBox

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False

try:
    from pyqtgraph.exporters.Exporter import Exporter
    from pyqtgraph.graphicsItems.PlotItem.PlotItem import PlotItem
    PG_EXPORT_OK = True
except ImportError:
    PG_EXPORT_OK = False


DATA_HEADERS = [
    "top_time",
    "top_pad_ua",
    "bottom_time",
    "bottom_mean_ua",
    "top_average_time",
    "top_average_ua",
    "bottom_average_time",
    "bottom_average_ua",
]


def _as_excel_time(timestamp):
    if timestamp is None:
        return None
    try:
        return datetime.fromtimestamp(float(timestamp))
    except (OSError, OverflowError, TypeError, ValueError):
        return None


def _write_time_cell(cell, timestamp):
    value = _as_excel_time(timestamp)
    if value is None:
        cell.value = None
        return
    cell.value = value
    cell.number_format = "yyyy-mm-dd hh:mm:ss.000"


class VisibleXlsxExporter(Exporter):
    Name = "Visible Viewer Data XLSX"

    def parameters(self):
        return None

    def export(self, fileName=None, toBytes=False, copy=False):
        if fileName is None:
            self.fileSaveDialog(filter=["*.xlsx"])
            return

        if not PG_EXPORT_OK or not isinstance(self.item, PlotItem):
            self._show_error("Visible XLSX export sadece plot alanlari icin kullanilabilir.")
            return

        provider = getattr(self.item, "_visible_xlsx_export_provider", None)
        if provider is None:
            self._show_error("Bu grafik icin Visible XLSX export verisi bulunamadi.")
            return

        if not OPENPYXL_OK:
            self._show_error("openpyxl kutuphanesi bulunamadi. XLSX yazilamiyor.")
            return

        if not fileName.lower().endswith(".xlsx"):
            fileName += ".xlsx"

        try:
            export_data = provider()
            self._write_workbook(fileName, export_data)
        except Exception as exc:
            self._show_error(f"Visible XLSX export hatasi:\n{type(exc).__name__}: {exc}")

    def _write_workbook(self, fileName, export_data):
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"

        ws.append(DATA_HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)

        series = export_data.get("series", {})
        max_len = max(
            len(series.get("top_pad", {}).get("x", [])),
            len(series.get("bottom_mean", {}).get("x", [])),
            len(series.get("top_average", {}).get("x", [])),
            len(series.get("bottom_average", {}).get("x", [])),
            0,
        )

        for index in range(max_len):
            row = ws.max_row + 1
            self._write_pair(ws, row, 1, series.get("top_pad", {}), index)
            self._write_pair(ws, row, 3, series.get("bottom_mean", {}), index)
            self._write_pair(ws, row, 5, series.get("top_average", {}), index)
            self._write_pair(ws, row, 7, series.get("bottom_average", {}), index)

        for column in ("A", "C", "E", "G"):
            ws.column_dimensions[column].width = 24
        for column in ("B", "D", "F", "H"):
            ws.column_dimensions[column].width = 16

        meta_ws = wb.create_sheet("Metadata")
        meta_ws.append(["key", "value"])
        for cell in meta_ws[1]:
            cell.font = Font(bold=True)
        for key, value in export_data.get("metadata", []):
            meta_ws.append([key, value])
            if isinstance(value, datetime):
                meta_ws.cell(meta_ws.max_row, 2).number_format = "yyyy-mm-dd hh:mm:ss.000"
        meta_ws.column_dimensions["A"].width = 30
        meta_ws.column_dimensions["B"].width = 32

        wb.save(fileName)

    @staticmethod
    def _write_pair(ws, row, column, series, index):
        times = series.get("x", [])
        values = series.get("y", [])
        if index >= len(times) or index >= len(values):
            return
        _write_time_cell(ws.cell(row, column), times[index])
        ws.cell(row, column + 1).value = values[index]
        ws.cell(row, column + 1).number_format = "0.000000"

    @staticmethod
    def _show_error(message: str):
        QMessageBox.warning(None, "Visible XLSX Export", message)


def register_visible_xlsx_exporter():
    if not PG_EXPORT_OK:
        return
    if VisibleXlsxExporter not in Exporter.Exporters:
        VisibleXlsxExporter.register()
