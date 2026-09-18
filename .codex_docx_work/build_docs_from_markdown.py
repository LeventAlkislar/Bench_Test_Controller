from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"C:\UNAM DRIVE\TUSEB\SGI Prototip Gelistirme\1B Sistem Muhendisligi\Bench Tests\CGM Control Software\Bench Test Controller")
REFERENCE = Path(r"C:\Users\LEVENT ALKISLAR\OneDrive\Desktop\Bilkent Üniversitesi UNAM.docx")
OUTPUT_DIR = ROOT / "docs" / "formatted"


def clear_body(doc: Document) -> None:
    body = doc._body._element
    sect_pr = None
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            sect_pr = child
        body.remove(child)
    if sect_pr is not None:
        body.append(sect_pr)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "D9D9D9", size: str = "6") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(table, top=80, start=100, bottom=80, end=100) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tbl_cell_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def format_paragraph(paragraph, *, justify=False, keep_with_next=False) -> None:
    pf = paragraph.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    if justify:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if keep_with_next:
        pf.keep_with_next = True


def add_text_with_inline_code(paragraph, text: str, *, bold=False) -> None:
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        else:
            run = paragraph.add_run(part)
            run.bold = bold


def clean_heading(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^\d+(\.\d+)*\.\s+", "", text)
    text = re.sub(r"^\d+(\.\d+)*\s+", "", text)
    return text


def is_ascii_table(block: list[str]) -> bool:
    return any(line.startswith("+") and line.endswith("+") for line in block) and any(
        line.startswith("|") for line in block
    )


def parse_ascii_table(block: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in block:
        if not line.startswith("|"):
            continue
        inner = line.strip()[1:-1]
        rows.append([cell.strip() for cell in inner.split("|")])
    return rows


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_cell_margins(table)
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx].cells
        for c_idx in range(cols):
            text = row[c_idx] if c_idx < len(row) else ""
            cell = cells[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_borders(cell)
            if r_idx == 0:
                set_cell_shading(cell, "D9EAF7")
            elif r_idx % 2 == 0:
                set_cell_shading(cell, "F7F7F7")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            add_text_with_inline_code(p, text, bold=(r_idx == 0))
            for run in p.runs:
                run.font.size = Pt(9)
                if r_idx == 0:
                    run.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def add_code_block(doc: Document, block: list[str]) -> None:
    if is_ascii_table(block):
        add_table(doc, parse_ascii_table(block))
        return
    for line in block:
        p = doc.add_paragraph(style="Normal")
        format_paragraph(p)
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)


def add_list_item(doc: Document, line: str) -> None:
    stripped = line.strip()
    ordered = re.match(r"^\d+\.\s+(.*)$", stripped)
    unordered = re.match(r"^[-*]\s+(.*)$", stripped)
    text = stripped
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    add_text_with_inline_code(p, text)


def add_paragraphs_from_buffer(doc: Document, buffer: list[str]) -> None:
    if not buffer:
        return
    text = " ".join(line.strip() for line in buffer if line.strip())
    if not text:
        return
    if text.endswith(":") and len(text) < 80:
        p = doc.add_paragraph(style="Heading 3")
        add_text_with_inline_code(p, clean_heading(text), bold=True)
        p.paragraph_format.keep_with_next = True
        return
    p = doc.add_paragraph(style="Normal")
    format_paragraph(p, justify=True)
    add_text_with_inline_code(p, text)


def build_doc(markdown_path: Path, output_path: Path, title: str) -> None:
    shutil.copy2(REFERENCE, output_path)
    doc = Document(str(output_path))
    clear_body(doc)

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(0.79)
    section.top_margin = Inches(1.30)
    section.bottom_margin = Inches(0.98)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(6)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size, color in [
        ("Title", 18, "000000"),
        ("Heading 1", 12, "262626"),
        ("Heading 2", 11, "262626"),
        ("Heading 3", 11, "0D0D0D"),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 6)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    title_p = doc.add_paragraph(style="Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(18)

    subtitle = doc.add_paragraph(style="Normal")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle_run = subtitle.add_run("Bench Test Controller")
    subtitle_run.bold = True

    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_block: list[str] = []
    para_buffer: list[str] = []
    skip_first_heading = True

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("```"):
            if in_code:
                add_paragraphs_from_buffer(doc, para_buffer)
                para_buffer = []
                add_code_block(doc, code_block)
                code_block = []
                in_code = False
            else:
                add_paragraphs_from_buffer(doc, para_buffer)
                para_buffer = []
                in_code = True
            continue
        if in_code:
            code_block.append(line)
            continue
        if not line.strip():
            add_paragraphs_from_buffer(doc, para_buffer)
            para_buffer = []
            continue
        if re.match(r"^(Version|Audience|Scope|Related document):", line.strip()):
            add_paragraphs_from_buffer(doc, para_buffer)
            para_buffer = []
            p = doc.add_paragraph(style="Normal")
            format_paragraph(p, justify=False)
            add_text_with_inline_code(p, line.strip())
            continue
        if line.startswith(">"):
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading:
            add_paragraphs_from_buffer(doc, para_buffer)
            para_buffer = []
            level = len(heading.group(1))
            text = clean_heading(heading.group(2))
            if skip_first_heading and level == 1:
                skip_first_heading = False
                continue
            style = "Heading 1" if level == 2 else "Heading 2" if level == 3 else "Heading 3"
            p = doc.add_paragraph(style=style)
            add_text_with_inline_code(p, text, bold=True)
            continue
        if re.match(r"^\s*(\d+\.|[-*])\s+", line):
            add_paragraphs_from_buffer(doc, para_buffer)
            para_buffer = []
            add_list_item(doc, line)
            continue
        para_buffer.append(line)

    add_paragraphs_from_buffer(doc, para_buffer)
    doc.save(str(output_path))


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    build_doc(
        ROOT / "docs" / "user_manual.md",
        OUTPUT_DIR / "Bench Test Controller User Manual formatted.docx",
        "Bench Test Controller User Manual",
    )
    build_doc(
        ROOT / "docs" / "standard_test_procedure.md",
        OUTPUT_DIR / "Bench Test Controller Standard Operating Procedure formatted.docx",
        "Bench Test Controller Standard Operating Procedure",
    )


if __name__ == "__main__":
    main()
